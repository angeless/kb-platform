"""Tests for the Word (.docx) parser."""

import io

import pytest
from docx import Document

from worker.parsers.word_parser import parse
from worker.parsers import get_parser, is_parseable


def _make_docx(paragraphs: list[tuple[str, str | None]]) -> bytes:
    """Create a .docx in memory.

    Each item is (text, style_name). style_name=None uses default (Normal).
    """
    doc = Document()
    for text, style in paragraphs:
        if style:
            doc.add_paragraph(text, style=style)
        else:
            doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


class TestWordParser:
    def test_basic_paragraphs(self):
        """Normal paragraphs should produce chunks."""
        docx_bytes = _make_docx([
            ("First paragraph.", None),
            ("Second paragraph.", None),
        ])
        result = parse(docx_bytes, "test.docx")
        assert len(result) >= 1
        assert "First paragraph" in result[0]["content_text"]
        assert result[0]["tags"]["source_type"] == "docx"
        assert result[0]["tags"]["filename"] == "test.docx"

    def test_headings_create_sections(self):
        """H1/H2 headings should split into separate chunks with heading_text."""
        docx_bytes = _make_docx([
            ("Chapter One", "Heading 1"),
            ("Content under chapter one.", None),
            ("Section A", "Heading 2"),
            ("Content under section A.", None),
        ])
        result = parse(docx_bytes, "headings.docx")
        assert len(result) >= 2

        # First chunk: Chapter One
        ch1 = result[0]
        assert "Chapter One" in ch1["content_text"]
        assert ch1["tags"]["heading_level"] == 1
        assert ch1["tags"]["heading_text"] == "Chapter One"

        # Second chunk: Section A
        sec_a = result[1]
        assert "Section A" in sec_a["content_text"]
        assert sec_a["tags"]["heading_level"] == 2
        assert sec_a["tags"]["heading_text"] == "Section A"

    def test_empty_document(self):
        """Empty .docx should return empty list, not error."""
        doc = Document()
        buf = io.BytesIO()
        doc.save(buf)
        result = parse(buf.getvalue(), "empty.docx")
        assert result == []

    def test_invalid_file_raises_valueerror(self):
        """Non-docx bytes should raise ValueError."""
        with pytest.raises(ValueError, match="不是有效的 .docx 文件"):
            parse(b"this is not a docx file", "bad.docx")

    def test_doc_legacy_returns_none(self):
        """Legacy .doc type should not have a parser (get_parser returns None)."""
        assert get_parser("doc") is None

    def test_docx_registered(self):
        """docx type should have a parser registered."""
        assert get_parser("docx") is not None
        assert is_parseable("docx") is True

    def test_page_or_timestamp_format(self):
        """Chunks should have para-N format timestamps."""
        docx_bytes = _make_docx([
            ("Only paragraph.", None),
        ])
        result = parse(docx_bytes, "ts.docx")
        assert len(result) == 1
        assert result[0]["page_or_timestamp"].startswith("para-")

    def test_heading_level_zero_for_normal(self):
        """Normal paragraphs should have heading_level=0."""
        docx_bytes = _make_docx([
            ("Just body text.", None),
        ])
        result = parse(docx_bytes, "body.docx")
        assert result[0]["tags"]["heading_level"] == 0
        assert result[0]["tags"]["heading_text"] is None
